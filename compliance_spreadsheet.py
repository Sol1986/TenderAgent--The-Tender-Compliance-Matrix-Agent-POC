"""Excel projection of the authoritative package-level compliance report."""

from __future__ import annotations

import json
import tempfile
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import TYPE_CHECKING, Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

if TYPE_CHECKING:
    from main import ComplianceReport, RequirementItem, UnresolvedIssue

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")
DISQUALIFYING_FILL = PatternFill("solid", fgColor="F4CCCC")
REVIEW_AND_DISQUALIFYING_FILL = PatternFill("solid", fgColor="F4B183")

MATRIX_HEADERS = (
    "Item",
    "Requirement",
    "Category",
    "M/R",
    "Section",
    "Page",
    "Source Document",
    "Required At",
    "Severity",
    "Human Review",
    "Issue",
)

DETAIL_HEADERS = (
    "Item",
    "Requirement ID",
    "Requirement",
    "Structured Parameters",
    "Source Document",
    "Section",
    "Page",
    "Evidence Text",
    "External Reference",
    "External Reference Retrieved",
    "Amendment Detected",
    "Amendment Details",
    "Ambiguity Detected",
    "Ambiguity Reason",
    "Contradiction Detected",
    "Contradiction Reason",
    "Human Review Reason",
)

ISSUE_HEADERS = (
    "Issue ID",
    "Requirement ID",
    "Issue Type",
    "Severity",
    "Description",
    "Source Document",
    "Section",
    "Page",
    "Required Action",
)

DOCUMENT_HEADERS = (
    "Document ID",
    "Filename",
    "Document Type",
    "Title",
    "Amendment Number",
    "Issue Date",
    "Processing Status",
    "Requirements Found",
)


def safe_excel_value(value: Any) -> Any:
    """Remove illegal controls and neutralize formula-like source text."""
    if value is None or not isinstance(value, str):
        return value
    cleaned = ILLEGAL_CHARACTERS_RE.sub("", value)
    if len(cleaned) > 32_767:
        cleaned = f"{cleaned[:32_755]} [truncated]"
    if cleaned.startswith(("=", "+", "-", "@")):
        return f"'{cleaned}"
    return cleaned


def display_enum(value: str | None) -> str:
    """Render internal enum names as concise spreadsheet labels."""
    if not value:
        return ""
    labels = {
        "BID_SUBMISSION": "Bid Submission",
        "BID_CLOSING": "Closing",
        "CONTRACT_AWARD": "Contract Award",
        "BEFORE_WORK_BEGINS": "Before Work Begins",
        "DURING_CONTRACT": "During Contract",
        "CONDITIONAL": "Conditional",
        "UNCLEAR": "Unclear",
    }
    return labels.get(value, value.replace("_", " ").title())


def unique_text(values: Iterable[str | None]) -> str:
    """Join non-empty values once while preserving their source order."""
    return "; ".join(
        dict.fromkeys(value.strip() for value in values if value and value.strip())
    )


def active_requirements(report: ComplianceReport) -> list[RequirementItem]:
    """Return the final active matrix population in stable report order."""
    return [
        item
        for item in report.requirements
        if item.analysis.is_active and item.analysis.version_status != "SUPERSEDED"
    ]


def issue_text(
    requirement: RequirementItem,
    linked_issues: Sequence[UnresolvedIssue],
) -> str:
    """Build a concise, source-derived issue label for the primary matrix."""
    return unique_text(
        [
            requirement.analysis.ambiguity_reason,
            requirement.analysis.contradiction_reason,
            *(issue.title for issue in linked_issues),
            requirement.analysis.review_reason,
        ]
    )


def write_sheet(
    worksheet: Worksheet,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    widths: dict[str, float],
    wrap_columns: set[int],
) -> None:
    """Write a filterable, frozen, consistently styled worksheet."""
    worksheet.append(list(headers))
    for row in rows:
        worksheet.append([safe_excel_value(value) for value in row])

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        f"A1:{worksheet.cell(worksheet.max_row, len(headers)).coordinate}"
    )
    worksheet.sheet_view.showGridLines = False

    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    for row in worksheet.iter_rows(min_row=2):
        for index, cell in enumerate(row, 1):
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=index in wrap_columns,
            )


def matrix_rows(report: ComplianceReport) -> list[tuple[Any, ...]]:
    """Create one concise row per active reconciled requirement."""
    issues_by_requirement: dict[str, list[UnresolvedIssue]] = {}
    for issue in report.unresolved_issues:
        for requirement_id in issue.affected_requirement_ids:
            issues_by_requirement.setdefault(requirement_id, []).append(issue)

    rows: list[tuple[Any, ...]] = []
    for item_number, item in enumerate(active_requirements(report), 1):
        primary = item.sources.evidence[0]
        rows.append(
            (
                item_number,
                item.requirement.text,
                display_enum(item.requirement.category),
                display_enum(item.requirement.requirement_type),
                primary.section or "",
                primary.page,
                primary.document_name,
                display_enum(item.requirement.required_at),
                display_enum(item.requirement.compliance_severity),
                "Yes" if item.analysis.requires_human_review else "No",
                issue_text(item, issues_by_requirement.get(item.item_id, [])),
            )
        )
    return rows


def detail_rows(report: ComplianceReport) -> list[tuple[Any, ...]]:
    """Repeat stable requirement IDs across every supporting evidence row."""
    rows: list[tuple[Any, ...]] = []
    for item_number, item in enumerate(active_requirements(report), 1):
        reference_names = unique_text(
            reference.reference_name for reference in item.sources.external_references
        )
        reference_statuses = unique_text(
            f"{reference.reference_name}: {'Yes' if reference.retrieved else 'No'}"
            for reference in item.sources.external_references
        )
        parameters = json.dumps(
            item.requirement.parameters,
            ensure_ascii=False,
            sort_keys=True,
        )
        for evidence in item.sources.evidence:
            rows.append(
                (
                    item_number,
                    item.item_id,
                    item.requirement.text,
                    parameters,
                    evidence.document_name,
                    evidence.section or "",
                    evidence.page,
                    evidence.text,
                    reference_names,
                    reference_statuses,
                    "Yes" if item.analysis.amendment_detected else "No",
                    item.analysis.amendment_details or "",
                    "Yes" if item.analysis.ambiguity_detected else "No",
                    item.analysis.ambiguity_reason or "",
                    "Yes" if item.analysis.contradiction_detected else "No",
                    item.analysis.contradiction_reason or "",
                    item.analysis.review_reason or "",
                )
            )
    return rows


def required_action(issue_type: str) -> str:
    """Map unresolved issue types to concrete proposal-team actions."""
    actions = {
        "AMBIGUOUS_REQUIREMENT": "Review the source and seek clarification if needed.",
        "CONTRADICTORY_REQUIREMENT": "Determine which clause controls before submission.",
        "MISSING_REFERENCED_DOCUMENT": "Obtain and review the referenced document.",
        "UNRESOLVED_CROSS_REFERENCE": "Identify and review the controlling reference.",
        "UNRESOLVED_AMENDMENT": "Confirm amendment precedence and current wording.",
        "PROCESSING_FAILURE": "Reprocess or manually inspect the source PDF.",
        "UNCLEAR_TIMING": "Confirm when compliance must be demonstrated.",
    }
    return actions.get(issue_type, "Review and resolve before submission.")


def issue_rows(report: ComplianceReport) -> list[tuple[Any, ...]]:
    """Project unresolved issues plus derived unclear-timing review items."""
    documents = {
        document.document_id: document for document in report.documents_analyzed
    }
    rows: list[tuple[Any, ...]] = []
    for issue in report.unresolved_issues:
        requirement_ids = issue.affected_requirement_ids or [""]
        source = documents.get(issue.source_document_id or "")
        for requirement_id in requirement_ids:
            rows.append(
                (
                    issue.issue_id,
                    requirement_id,
                    display_enum(issue.issue_type),
                    display_enum(issue.severity),
                    issue.description,
                    source.filename if source is not None else "",
                    issue.source_section or "",
                    issue.source_page,
                    required_action(issue.issue_type),
                )
            )

    for item in active_requirements(report):
        if item.requirement.required_at != "UNCLEAR":
            continue
        primary = item.sources.evidence[0]
        rows.append(
            (
                f"REVIEW-{item.item_id}-TIMING",
                item.item_id,
                "Unclear Timing",
                display_enum(item.requirement.compliance_severity),
                item.analysis.review_reason
                or "The package does not clearly establish the compliance timing.",
                primary.document_name,
                primary.section or "",
                primary.page,
                required_action("UNCLEAR_TIMING"),
            )
        )
    return rows


def document_rows(report: ComplianceReport) -> list[tuple[Any, ...]]:
    """List every registered PDF and its traceable final requirement count."""
    requirement_ids_by_document: dict[str, set[str]] = {}
    for item in report.requirements:
        for evidence in item.sources.evidence:
            requirement_ids_by_document.setdefault(evidence.document_id, set()).add(
                item.item_id
            )
    return [
        (
            document.document_id,
            document.filename,
            display_enum(document.document_type),
            document.title or "",
            document.amendment_number or "",
            document.issue_date or "",
            display_enum(document.processing_status),
            len(requirement_ids_by_document.get(document.document_id, set())),
        )
        for document in report.documents_analyzed
    ]


def apply_matrix_highlights(worksheet: Worksheet) -> None:
    """Use restrained row fills for human-review and disqualifying requirements."""
    for row in worksheet.iter_rows(min_row=2, max_col=len(MATRIX_HEADERS)):
        severity = row[8].value
        human_review = row[9].value
        if severity == "Disqualifying" and human_review == "Yes":
            fill = REVIEW_AND_DISQUALIFYING_FILL
        elif severity == "Disqualifying":
            fill = DISQUALIFYING_FILL
        elif human_review == "Yes":
            fill = REVIEW_FILL
        else:
            continue
        for cell in row:
            cell.fill = fill


def verify_workbook(path: Path) -> None:
    """Reopen the file and verify sheets, filters, frozen rows, and formula safety."""
    expected_sheets = [
        "Compliance Matrix",
        "Requirement Details",
        "Issues - Human Review",
        "Document Register",
    ]
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if workbook.sheetnames != expected_sheets:
            raise ValueError("Compliance workbook sheet structure is invalid")
        for worksheet in workbook.worksheets:
            if worksheet.freeze_panes != "A2" or not worksheet.auto_filter.ref:
                raise ValueError(
                    f"Worksheet formatting is incomplete: {worksheet.title}"
                )
            for row in worksheet.iter_rows():
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError("Compliance workbook must not contain formulas")
    finally:
        workbook.close()


def export_compliance_workbook(
    report: ComplianceReport,
    output_path: str | Path,
) -> Path:
    """Create and atomically save one four-sheet package compliance workbook."""
    target = Path(output_path).resolve()
    if target.suffix.casefold() != ".xlsx":
        raise ValueError("Compliance workbook output must use the .xlsx extension")
    target.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    matrix = workbook.active
    matrix.title = "Compliance Matrix"
    details = workbook.create_sheet("Requirement Details")
    issues = workbook.create_sheet("Issues - Human Review")
    documents = workbook.create_sheet("Document Register")

    write_sheet(
        matrix,
        MATRIX_HEADERS,
        matrix_rows(report),
        {
            "A": 8,
            "B": 55,
            "C": 22,
            "D": 16,
            "E": 18,
            "F": 8,
            "G": 30,
            "H": 20,
            "I": 16,
            "J": 15,
            "K": 45,
        },
        {2, 11},
    )
    apply_matrix_highlights(matrix)
    write_sheet(
        details,
        DETAIL_HEADERS,
        detail_rows(report),
        {
            "A": 8,
            "B": 18,
            "C": 50,
            "D": 35,
            "E": 30,
            "F": 18,
            "G": 8,
            "H": 65,
            "I": 35,
            "J": 35,
            "K": 18,
            "L": 45,
            "M": 18,
            "N": 45,
            "O": 20,
            "P": 45,
            "Q": 45,
        },
        {3, 4, 8, 9, 10, 12, 14, 16, 17},
    )
    write_sheet(
        issues,
        ISSUE_HEADERS,
        issue_rows(report),
        {
            "A": 24,
            "B": 18,
            "C": 28,
            "D": 16,
            "E": 60,
            "F": 30,
            "G": 18,
            "H": 8,
            "I": 50,
        },
        {5, 9},
    )
    write_sheet(
        documents,
        DOCUMENT_HEADERS,
        document_rows(report),
        {
            "A": 20,
            "B": 32,
            "C": 28,
            "D": 40,
            "E": 20,
            "F": 16,
            "G": 20,
            "H": 20,
        },
        {4},
    )

    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=target.parent,
            prefix=f".{target.name}.",
            suffix=".tmp.xlsx",
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        workbook.close()
        verify_workbook(temporary_path)
        temporary_path.replace(target)
    except Exception:
        workbook.close()
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        raise
    return target
