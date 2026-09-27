"""Offline end-to-end checks for one reconciled solicitation package."""

import json
from contextvars import ContextVar
from pathlib import Path
from threading import Lock
from typing import Any

from openpyxl import load_workbook

from main import (
    AmendmentResolutionDraft,
    ChunkFindings,
    ComplianceReport,
    ComplianceReportDraft,
    ExternalReferenceDraft,
    FindingRelationshipDecision,
    PackageResolutionDraft,
    ReconciliationDraft,
    ReconciliationGroupDraft,
    Requirement,
    RequirementEnrichment,
    RequirementParameter,
    run_compliance_solicitation_package,
)


class ParsedDocument:
    """Provide deterministic Markdown without invoking Docling."""

    def __init__(self, filename: str) -> None:
        self.filename = filename
        self.tables: list[Any] = []

    def export_to_markdown(self) -> str:
        """Return one source section for each registered PDF."""
        return f"## Requirements\nEvidence from {self.filename}."


class PackageModel:
    """Return scenario-specific structured outputs for every pipeline stage."""

    def __init__(
        self,
        findings: dict[str, list[Requirement]],
        reconciliation: ReconciliationDraft,
        resolution: PackageResolutionDraft,
        report: ComplianceReportDraft,
    ) -> None:
        self.findings = findings
        self.reconciliation = reconciliation
        self.resolution = resolution
        self.report = report
        self.calls: list[str] = []
        self.lock = Lock()

    def with_structured_output(self, schema: type) -> Any:
        """Mimic the provider's structured-output adapter."""
        model = self

        class Runnable:
            def invoke(self, prompt: str) -> Any:
                with model.lock:
                    model.calls.append(schema.__name__)
                if schema is ChunkFindings:
                    filename = next(name for name in model.findings if name in prompt)
                    return ChunkFindings(requirements=model.findings[filename])
                if schema is ReconciliationDraft:
                    return model.reconciliation
                if schema is PackageResolutionDraft:
                    return model.resolution
                if schema is ComplianceReportDraft:
                    return model.report
                raise AssertionError(f"Unexpected schema invocation: {schema}")

        return Runnable()


def requirement(
    text: str,
    evidence: str,
    *,
    status: str = "FOUND",
    category: str = "required_documents",
) -> Requirement:
    """Build one source-grounded extraction result."""
    return Requirement(
        category=category,
        requirement=text,
        status=status,
        source_section="Requirements",
        evidence=evidence,
    )


def package_files(package: Path, *filenames: str) -> None:
    """Create discoverable placeholders because parsing is injected."""
    package.mkdir()
    for filename in filenames:
        (package / filename).write_bytes(b"%PDF-test-placeholder")


def run_package(
    tmp_path: Path,
    filenames: tuple[str, ...],
    model: PackageModel,
    parser: Any | None = None,
) -> tuple[ComplianceReport, Path, Path]:
    """Run every stage and return the authoritative and projected artifacts."""
    package = tmp_path / "tender_package"
    package_files(package, *filenames)
    json_path = tmp_path / "outputs" / "compliance_report.json"
    workbook_path = tmp_path / "outputs" / "compliance_matrix.xlsx"
    report = run_compliance_solicitation_package(
        model=model,
        package_path=package,
        output_path=json_path,
        spreadsheet_output_path=workbook_path,
        parser=parser or (lambda path: ParsedDocument(path.name)),
    )
    return report, json_path, workbook_path


def test_package_pipeline_applies_amendment_and_links_uploaded_reference(
    tmp_path: Path,
) -> None:
    """A controlling amendment becomes the single matrix row with all evidence."""
    findings = {
        "amendment_001.pdf": [
            requirement(
                "Provide bid security equal to 10% of the bid price.",
                "Clause 2 is replaced: provide bid security of 10%.",
                category="bonding_security",
            )
        ],
        "main_solicitation.pdf": [
            requirement(
                "Provide bid security equal to 5% under Standard Instructions 2003.",
                "Provide 5% bid security in accordance with Standard Instructions 2003.",
                category="bonding_security",
            )
        ],
        "standard_instructions_2003.pdf": [
            requirement(
                "Submit the prescribed bid-security form.",
                "Bid security must use the prescribed form.",
                category="bonding_security",
            )
        ],
    }
    reconciliation = ReconciliationDraft(
        groups=[
            ReconciliationGroupDraft(
                candidate_ids=["CAND-0001", "CAND-0002", "CAND-0003"],
                requirement="Provide 10% bid security using the prescribed form.",
                category="bonding_security",
                extraction_status="FOUND",
                relationships=[
                    FindingRelationshipDecision(
                        left_candidate_id="CAND-0001",
                        right_candidate_id="CAND-0002",
                        relationship="SUPERSEDES",
                        reason="Amendment 001 replaces the original percentage.",
                    ),
                    FindingRelationshipDecision(
                        left_candidate_id="CAND-0002",
                        right_candidate_id="CAND-0003",
                        relationship="SUPPLEMENTS",
                        reason="The instructions supply the required form.",
                    ),
                ],
                requires_human_review=True,
                review_reason="Apply the amendment before using the requirement.",
            )
        ]
    )
    resolution = PackageResolutionDraft(
        amendments=[
            AmendmentResolutionDraft(
                requirement_id="REQ-0001",
                action="REPLACE",
                superseding_candidate_id="CAND-0001",
                superseded_candidate_ids=["CAND-0002"],
                current_requirement="Provide bid security equal to 10% of the bid price using the prescribed form.",
                reason="Amendment 001 explicitly replaces the original amount.",
            )
        ],
        external_references=[
            ExternalReferenceDraft(
                requirement_id="REQ-0001",
                reference_name="Standard Instructions 2003",
            )
        ],
    )
    report_draft = ComplianceReportDraft(
        requirements=[
            RequirementEnrichment(
                requirement_id="REQ-0001",
                requirement_type="MANDATORY",
                required_at="BID_SUBMISSION",
                compliance_severity="DISQUALIFYING",
                parameters=[
                    RequirementParameter(name="bid_security", value="10% of bid price")
                ],
            )
        ]
    )
    report, json_path, workbook_path = run_package(
        tmp_path,
        (
            "amendment_001.pdf",
            "main_solicitation.pdf",
            "standard_instructions_2003.pdf",
        ),
        PackageModel(findings, reconciliation, resolution, report_draft),
    )

    assert json.loads(json_path.read_text(encoding="utf-8")) == report.model_dump(
        mode="json"
    )
    assert len(report.requirements) == 1
    item = report.requirements[0]
    assert "10%" in item.requirement.text
    assert item.analysis.version_status == "CURRENT"
    assert len(item.sources.evidence) == 3
    assert item.sources.external_references[0].retrieved is True
    assert item.sources.external_references[0].matched_document_id

    workbook = load_workbook(workbook_path, read_only=True)
    try:
        assert workbook["Compliance Matrix"].max_row == 2
        assert workbook["Requirement Details"].max_row == 4
        assert workbook["Document Register"].max_row == 4
    finally:
        workbook.close()


def test_package_pipeline_keeps_two_missing_references_visible(tmp_path: Path) -> None:
    """Absent incorporated documents produce separate material review issues."""
    findings = {
        "annex_a.pdf": [],
        "main_solicitation.pdf": [
            requirement(
                "Comply with Annex E.",
                "The bidder must comply with Annex E.",
                status="EXTERNAL_REFERENCE",
            ),
            requirement(
                "Comply with Standard Instructions 2003.",
                "The bidder must comply with Standard Instructions 2003.",
                status="EXTERNAL_REFERENCE",
            ),
        ],
    }
    reconciliation = ReconciliationDraft(
        groups=[
            ReconciliationGroupDraft(
                candidate_ids=[candidate_id],
                requirement=text,
                category="required_documents",
                extraction_status="EXTERNAL_REFERENCE",
            )
            for candidate_id, text in (
                ("CAND-0001", "Comply with Annex E."),
                ("CAND-0002", "Comply with Standard Instructions 2003."),
            )
        ]
    )
    resolution = PackageResolutionDraft(
        external_references=[
            ExternalReferenceDraft(requirement_id="REQ-0001", reference_name="Annex E"),
            ExternalReferenceDraft(
                requirement_id="REQ-0002",
                reference_name="Standard Instructions 2003",
            ),
        ]
    )
    report_draft = ComplianceReportDraft(
        requirements=[
            RequirementEnrichment(
                requirement_id=f"REQ-{number:04d}",
                requirement_type="REQUIRED",
                required_at="UNCLEAR",
                compliance_severity="MAJOR",
            )
            for number in (1, 2)
        ]
    )
    report, _, workbook_path = run_package(
        tmp_path,
        ("annex_a.pdf", "main_solicitation.pdf"),
        PackageModel(findings, reconciliation, resolution, report_draft),
    )

    missing = [
        issue
        for issue in report.unresolved_issues
        if issue.issue_type == "MISSING_REFERENCED_DOCUMENT"
    ]
    assert len(missing) == 2
    assert report.summary.unresolved_external_references == 2
    assert all(item.analysis.requires_human_review for item in report.requirements)

    workbook = load_workbook(workbook_path, read_only=True)
    try:
        assert workbook["Compliance Matrix"].max_row == 3
        assert workbook["Issues - Human Review"].max_row >= 3
    finally:
        workbook.close()


def test_package_pipeline_preserves_contradictory_timing(tmp_path: Path) -> None:
    """Incompatible timing remains unresolved instead of being silently merged."""
    findings = {
        "annex_a.pdf": [
            requirement(
                "Provide security clearances with the bid.",
                "Clearance proof must accompany the bid.",
                category="security",
            )
        ],
        "main_solicitation.pdf": [
            requirement(
                "Provide security clearances before contract award.",
                "Clearance proof is required before contract award.",
                category="security",
            )
        ],
    }
    reconciliation = ReconciliationDraft(
        groups=[
            ReconciliationGroupDraft(
                candidate_ids=["CAND-0001", "CAND-0002"],
                requirement="Provide security clearance evidence at an unresolved time.",
                category="security",
                extraction_status="FOUND",
                relationships=[
                    FindingRelationshipDecision(
                        left_candidate_id="CAND-0001",
                        right_candidate_id="CAND-0002",
                        relationship="CONTRADICTS",
                        reason="The documents give different compliance times.",
                    )
                ],
                ambiguity_detected=True,
                ambiguity_reason="The controlling submission time is unclear.",
                contradiction_detected=True,
                contradiction_reason="One source requires bid submission and another contract award.",
                requires_human_review=True,
                review_reason="Confirm the controlling timing with the buyer.",
            )
        ]
    )
    report_draft = ComplianceReportDraft(
        requirements=[
            RequirementEnrichment(
                requirement_id="REQ-0001",
                requirement_type="MANDATORY",
                required_at="UNCLEAR",
                compliance_severity="DISQUALIFYING",
            )
        ]
    )
    report, _, workbook_path = run_package(
        tmp_path,
        ("annex_a.pdf", "main_solicitation.pdf"),
        PackageModel(
            findings,
            reconciliation,
            PackageResolutionDraft(),
            report_draft,
        ),
    )

    item = report.requirements[0]
    assert item.analysis.ambiguity_detected is True
    assert item.analysis.contradiction_detected is True
    assert item.analysis.requires_human_review is True
    issue_types = {issue.issue_type for issue in report.unresolved_issues}
    assert "AMBIGUOUS_REQUIREMENT" in issue_types
    assert "CONTRADICTORY_REQUIREMENT" in issue_types

    workbook = load_workbook(workbook_path, read_only=False)
    try:
        matrix = workbook["Compliance Matrix"]
        assert matrix["J2"].value == "Yes"
        assert "tim" in matrix["K2"].value.casefold()
    finally:
        workbook.close()


def test_parallel_document_work_preserves_parent_context(tmp_path: Path) -> None:
    """Propagate the active parent context into document worker threads."""
    findings = {
        "main_solicitation.pdf": [
            requirement(
                "Submit the signed offer.",
                "The bidder must submit the signed offer.",
                category="submission",
            )
        ]
    }
    model = PackageModel(
        findings,
        ReconciliationDraft(
            groups=[
                ReconciliationGroupDraft(
                    candidate_ids=["CAND-0001"],
                    requirement="Submit the signed offer.",
                    category="submission",
                    extraction_status="FOUND",
                )
            ]
        ),
        PackageResolutionDraft(),
        ComplianceReportDraft(
            requirements=[
                RequirementEnrichment(
                    requirement_id="REQ-0001",
                    requirement_type="MANDATORY",
                    required_at="BID_SUBMISSION",
                    compliance_severity="DISQUALIFYING",
                )
            ]
        ),
    )
    trace_marker: ContextVar[str] = ContextVar("trace_marker")
    parser_markers: list[str] = []

    def traced_parser(path: Path) -> ParsedDocument:
        parser_markers.append(trace_marker.get())
        return ParsedDocument(path.name)

    token = trace_marker.set("package-parent")
    try:
        run_package(
            tmp_path,
            ("main_solicitation.pdf",),
            model,
            parser=traced_parser,
        )
    finally:
        trace_marker.reset(token)

    assert parser_markers == ["package-parent"]
