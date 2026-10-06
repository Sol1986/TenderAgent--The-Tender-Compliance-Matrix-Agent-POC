"""Offline validation for the Phase E package-level compliance report."""

import json
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app import main
from app.compliance_spreadsheet import (
    export_compliance_workbook,
    issue_text,
    matrix_rows,
)
from app.main import (
    DEFAULT_COMPLIANCE_MATRIX_PATH,
    DEFAULT_COMPLIANCE_REPORT_PATH,
    ComplianceReportDraft,
    DocumentExtractionResult,
    Evidence,
    ExternalReference,
    PackageResolutionResult,
    ReconciledRequirement,
    RequirementEnrichment,
    RequirementParameter,
    UnresolvedIssue,
    build_compliance_report,
    export_compliance_matrix_workbook,
    project_compliance_matrix,
    save_compliance_report,
)
from app.solicitation_package import SolicitationDocument


class ReportModel:
    """Return supplied report enrichments without external model calls."""

    def __init__(self, draft: ComplianceReportDraft) -> None:
        self.draft = draft
        self.prompt = ""

    def with_structured_output(self, schema: type) -> Any:
        model = self

        class Runnable:
            def invoke(self, prompt: str) -> ComplianceReportDraft:
                if schema is not ComplianceReportDraft:
                    raise AssertionError(f"Unexpected schema invocation: {schema}")
                model.prompt = prompt
                return model.draft

        return Runnable()


def source_document(
    number: int,
    filename: str,
    document_type: str,
    *,
    status: str = "COMPLETE",
    error: str | None = None,
    solicitation_number: str | None = None,
    title: str | None = None,
) -> SolicitationDocument:
    """Build predictable document registry entries."""
    return SolicitationDocument(
        document_id=f"DOC-{number:012X}",
        filename=filename,
        document_type=document_type,
        solicitation_number=solicitation_number,
        title=title,
        processing_status=status,
        processing_error=error,
    )


def requirement(
    number: int,
    text: str,
    source: SolicitationDocument,
    *,
    ambiguity: bool = False,
    amendment: bool = False,
    version_status: str = "UNCHANGED",
    active: bool = True,
    references: list[ExternalReference] | None = None,
) -> ReconciledRequirement:
    """Build one Phase D requirement with trusted evidence."""
    return ReconciledRequirement(
        requirement_id=f"REQ-{number:04d}",
        requirement=text,
        category="bonding_security",
        extraction_status="EXTERNAL_REFERENCE" if references else "FOUND",
        candidate_ids=[f"CAND-{number:04d}"],
        active_candidate_ids=[f"CAND-{number:04d}"],
        evidence=[
            Evidence(
                document_id=source.document_id,
                document_name=source.filename,
                section=f"Section {number}",
                page=number,
                text=f"Evidence for requirement {number}.",
            )
        ],
        external_references=references or [],
        amendment_detected=amendment,
        amendment_details="An amendment changed this obligation."
        if amendment
        else None,
        version_status=version_status,
        is_active=active,
        ambiguity_detected=ambiguity,
        ambiguity_reason="The required timing is unclear." if ambiguity else None,
        requires_human_review=ambiguity,
        review_reason="Confirm the controlling timing." if ambiguity else None,
    )


def report_resolution() -> tuple[PackageResolutionResult, ComplianceReportDraft]:
    """Create a package with current, unresolved, and deleted requirements."""
    main = source_document(
        1,
        "main_solicitation.pdf",
        "MAIN_SOLICITATION",
        solicitation_number="SOL-2026-17",
        title="Bridge Engineering Services",
    )
    amendment = source_document(2, "amendment_001.pdf", "AMENDMENT")
    failed = source_document(
        3,
        "Annex_E.pdf",
        "ANNEX",
        status="FAILED",
        error="The annex could not be parsed.",
    )
    missing_reference = ExternalReference(
        reference_name="Standard Instructions 2003",
        referenced_from_document_id=main.document_id,
        referenced_from_section="Section 2",
        retrieved=False,
    )
    requirements = [
        requirement(
            1,
            "Provide bid security equal to 10% of the bid price.",
            amendment,
            amendment=True,
            version_status="CURRENT",
        ),
        requirement(
            2,
            "Hold the required security clearance.",
            main,
            ambiguity=True,
            references=[missing_reference],
        ),
        requirement(
            3,
            "Attend the deleted site meeting.",
            main,
            amendment=True,
            version_status="DELETED",
            active=False,
        ),
    ]
    issues = [
        UnresolvedIssue(
            issue_id="ISSUE-0001",
            issue_type="MISSING_REFERENCED_DOCUMENT",
            title="Standard Instructions 2003 was not uploaded",
            description="REQ-0002 depends on a document absent from the package.",
            affected_requirement_ids=["REQ-0002"],
            source_document_id=main.document_id,
            source_section="Section 2",
        ),
        UnresolvedIssue(
            issue_id="ISSUE-0002",
            issue_type="PROCESSING_FAILURE",
            title="Annex_E.pdf could not be processed",
            description="The annex could not be parsed.",
            severity="HIGH",
            source_document_id=failed.document_id,
        ),
    ]
    resolution = PackageResolutionResult(
        documents=[main, amendment, failed],
        document_results=[
            DocumentExtractionResult(document=main, text_chunks=1),
            DocumentExtractionResult(document=amendment, text_chunks=1),
            DocumentExtractionResult(document=failed),
        ],
        candidate_requirements=[],
        reconciled_requirements=requirements,
        external_references=[missing_reference],
        unresolved_issues=issues,
    )
    draft = ComplianceReportDraft(
        requirements=[
            RequirementEnrichment(
                requirement_id="REQ-0001",
                requirement_type="MANDATORY",
                required_at="BID_SUBMISSION",
                compliance_severity="DISQUALIFYING",
                consequence="The submission may be rejected without bid security.",
                parameters=[
                    RequirementParameter(
                        name="bid_security_amount", value="10% of bid price"
                    )
                ],
            ),
            RequirementEnrichment(
                requirement_id="REQ-0002",
                requirement_type="REQUIRED",
                required_at="UNCLEAR",
                compliance_severity="UNKNOWN",
            ),
            RequirementEnrichment(
                requirement_id="REQ-0003",
                requirement_type="MANDATORY",
                required_at="BID_SUBMISSION",
                compliance_severity="MAJOR",
            ),
        ]
    )
    return resolution, draft


def test_final_report_counts_only_active_requirements_and_projects_matrix() -> None:
    """Deleted rows stay traceable but never inflate active summary or matrix counts."""
    resolution, draft = report_resolution()
    model = ReportModel(draft)

    report = build_compliance_report(resolution, model)

    assert report.solicitation_number == "SOL-2026-17"
    assert report.solicitation_title == "Bridge Engineering Services"
    assert report.summary.model_dump() == {
        "documents_processed": 2,
        "documents_failed": 1,
        "total_requirements": 2,
        "mandatory_requirements": 1,
        "disqualifying_requirements": 1,
        "ambiguous_requirements": 1,
        "contradictory_requirements": 0,
        "external_references_found": 1,
        "unresolved_external_references": 1,
        "amendments_found": 1,
        "unresolved_amendments": 0,
        "human_review_required": 1,
    }
    assert report.requirements[2].analysis.version_status == "SUPERSEDED"
    matrix = project_compliance_matrix(report)
    assert [row.item for row in matrix] == ["REQ-0001", "REQ-0002"]
    assert matrix[1].missing_reference is True
    assert "bid/no-bid verdict" in model.prompt


def test_client_matrix_keeps_only_obligations_and_flags_missing_source() -> None:
    """Avoid falsely labeling context as Required or inventing source coordinates."""
    resolution, draft = report_resolution()
    report = build_compliance_report(resolution, ReportModel(draft))
    report.requirements[0].requirement.requirement_type = "INFORMATIONAL"
    conditional = report.requirements[1]
    conditional.requirement.requirement_type = "CONDITIONAL"
    conditional.requirement.section = None
    conditional.requirement.page = None
    conditional.sources.evidence[0].section = None
    conditional.sources.evidence[0].page = None

    rows = matrix_rows(report)

    assert len(rows) == 1
    assert rows[0][3] == "Required"
    assert rows[0][2] == ""
    assert rows[0][6] is None
    assert "Conditional requirement" in rows[0][7]
    assert "Source section not identified" in rows[0][7]
    assert "Source page not identified" in rows[0][7]


def test_matrix_comments_hide_internal_reconciliation_note(tmp_path: Path) -> None:
    """Keep useful review guidance without exposing fallback implementation text."""
    resolution, draft = report_resolution()
    report = build_compliance_report(resolution, ReportModel(draft))
    item = report.requirements[1]
    internal_note = (
        "Automated package reconciliation failed validation, so this candidate "
        "set was retained conservatively for human review."
    )
    item.analysis.ambiguity_reason = internal_note
    item.analysis.review_reason = f"{internal_note}; Confirm the controlling timing."
    report.unresolved_issues[0].description = internal_note

    comment = issue_text(item, [])

    assert internal_note not in comment
    assert comment == "Confirm the controlling timing."
    assert report.requirements[1].analysis.ambiguity_reason == internal_note

    workbook_path = export_compliance_workbook(report, tmp_path / "matrix.xlsx")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        assert all(
            internal_note not in cell.value
            for sheet in workbook
            for row in sheet
            for cell in row
            if isinstance(cell.value, str)
        )
    finally:
        workbook.close()


@pytest.mark.parametrize("mode", ["missing", "duplicate", "unknown"])
def test_enrichment_must_cover_every_requirement_exactly_once(mode: str) -> None:
    """The model cannot drop, duplicate, or add final requirements."""
    resolution, complete = report_resolution()
    enrichments = list(complete.requirements)
    if mode == "missing":
        enrichments.pop()
    elif mode == "duplicate":
        enrichments.append(enrichments[0])
    else:
        enrichments[-1] = enrichments[-1].model_copy(
            update={"requirement_id": "REQ-9999"}
        )

    with pytest.raises(ValueError, match="enrichment"):
        build_compliance_report(
            resolution,
            ReportModel(ComplianceReportDraft(requirements=enrichments)),
        )


def test_large_final_report_is_enriched_in_complete_batches() -> None:
    """A provider response limited to one batch cannot silently lose later rows."""
    resolution, _ = report_resolution()
    source = resolution.documents[0]
    resolution.reconciled_requirements.extend(
        requirement(number, f"Submit document {number}.", source)
        for number in range(4, 47)
    )

    class BoundedReportModel:
        """Return only the requirements present in each individual prompt."""

        def __init__(self) -> None:
            self.calls: list[list[str]] = []

        def with_structured_output(self, schema: type) -> Any:
            model = self

            class Runnable:
                def invoke(self, prompt: str) -> ComplianceReportDraft:
                    if schema is not ComplianceReportDraft:
                        raise AssertionError(f"Unexpected schema invocation: {schema}")
                    block = prompt.split("RESOLVED REQUIREMENTS:\n", 1)[1]
                    block = block.split("\n\nUNRESOLVED ISSUES:", 1)[0]
                    ids = [
                        json.loads(line)["requirement_id"]
                        for line in block.splitlines()
                        if line.startswith("{")
                    ]
                    model.calls.append(ids)
                    return ComplianceReportDraft(
                        requirements=[
                            RequirementEnrichment(
                                requirement_id=item_id,
                                requirement_type="REQUIRED",
                                required_at="BID_SUBMISSION",
                            )
                            for item_id in ids
                        ]
                    )

            return Runnable()

    model = BoundedReportModel()
    report = build_compliance_report(resolution, model)

    assert len(report.requirements) == 46
    assert [len(call) for call in model.calls] == [20, 20, 6]
    assert {item.item_id for item in report.requirements} == {
        f"REQ-{number:04d}" for number in range(1, 47)
    }


def test_final_report_retries_only_missing_enrichments() -> None:
    """A partial batch gets one focused repair without discarding valid rows."""
    resolution, complete = report_resolution()

    class PartialReportModel:
        def __init__(self) -> None:
            self.calls = 0

        def with_structured_output(self, schema: type) -> Any:
            model = self

            class Runnable:
                def invoke(self, prompt: str) -> ComplianceReportDraft:
                    if schema is not ComplianceReportDraft:
                        raise AssertionError(f"Unexpected schema invocation: {schema}")
                    model.calls += 1
                    if model.calls == 1:
                        return ComplianceReportDraft(
                            requirements=complete.requirements[:2]
                        )
                    assert "REQ-0003" in prompt
                    assert '"requirement_id":"REQ-0001"' not in prompt
                    return ComplianceReportDraft(requirements=complete.requirements[2:])

            return Runnable()

    model = PartialReportModel()
    report = build_compliance_report(resolution, model)

    assert model.calls == 2
    assert len(report.requirements) == 3


def test_parameter_floats_are_rejected_for_financial_precision() -> None:
    """Model parameters cannot introduce imprecise floating-point money values."""
    with pytest.raises(ValidationError, match="valid string"):
        RequirementParameter.model_validate(
            {"name": "minimum_insurance", "value": 5000000.50}
        )


@pytest.mark.parametrize("name", ["bid_security_amount", " bid_security_amount ", " "])
def test_parameter_names_cannot_silently_overwrite_values(
    name: str, tmp_path: Path
) -> None:
    """Preserve conflicting or unlabeled values and make them reviewable."""
    resolution, draft = report_resolution()
    draft.requirements[0].parameters.append(
        RequirementParameter(name=name, value="A conflicting source value")
    )

    report = build_compliance_report(resolution, ReportModel(draft))
    first = report.requirements[0]
    assert first.requirement.parameters["bid_security_amount"] == "10% of bid price"
    assert "A conflicting source value" in first.requirement.parameters.values()
    assert len(first.requirement.parameters) == 2
    assert first.analysis.requires_human_review is True
    assert first.analysis.ambiguity_detected is True
    assert any(
        issue.issue_type == "AMBIGUOUS_REQUIREMENT"
        and issue.affected_requirement_ids == ["REQ-0001"]
        and "parameter labels" in issue.title
        for issue in report.unresolved_issues
    )
    workbook_path = tmp_path / "matrix.xlsx"
    export_compliance_matrix_workbook(report, workbook_path)
    workbook = load_workbook(workbook_path, read_only=True)
    try:
        issue_rows = list(workbook["Issues - Human Review"].values)
        assert any(
            row[1] == "REQ-0001" and "Verify the labeled values" in str(row[4])
            for row in issue_rows[1:]
        )
    finally:
        workbook.close()


def test_identical_parameter_duplicates_collapse_without_review() -> None:
    """Repeating the same label and value adds no ambiguity or extra field."""
    resolution, draft = report_resolution()
    draft.requirements[0].parameters.append(
        RequirementParameter(name=" bid_security_amount ", value="10% of bid price")
    )

    report = build_compliance_report(resolution, ReportModel(draft))

    first = report.requirements[0]
    assert first.requirement.parameters == {"bid_security_amount": "10% of bid price"}
    assert first.analysis.requires_human_review is False
    assert not any(
        "parameter labels" in issue.title for issue in report.unresolved_issues
    )


def test_report_is_saved_as_valid_atomic_json(tmp_path: Path) -> None:
    """The persisted artifact round-trips through the Pydantic report schema."""
    resolution, draft = report_resolution()
    report = build_compliance_report(resolution, ReportModel(draft))
    output = tmp_path / "nested" / "compliance_report.json"

    saved_path = save_compliance_report(report, output)

    assert saved_path == output.resolve()
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["requirements"][0]["requirement"]["parameters"] == {
        "bid_security_amount": "10% of bid price"
    }
    assert payload == report.model_dump(mode="json")
    assert output.read_bytes().endswith(b"\n")
    assert list(output.parent.glob("*.tmp")) == []


def test_save_rejects_tampered_summary(tmp_path: Path) -> None:
    """Direct callers cannot persist report counts that disagree with its rows."""
    resolution, draft = report_resolution()
    report = build_compliance_report(resolution, ReportModel(draft))
    report.summary.total_requirements = 999

    with pytest.raises(ValueError, match="summary"):
        save_compliance_report(report, tmp_path / "invalid.json")

    assert not (tmp_path / "invalid.json").exists()


def test_empty_package_report_skips_model_and_preserves_failure() -> None:
    """A failed package still produces a truthful validated report with no rows."""
    failed = source_document(
        1,
        "broken.pdf",
        "OTHER",
        status="FAILED",
        error="Unable to parse PDF.",
    )
    issue = UnresolvedIssue(
        issue_id="ISSUE-0001",
        issue_type="PROCESSING_FAILURE",
        title="broken.pdf could not be processed",
        description="Unable to parse PDF.",
        severity="HIGH",
        source_document_id=failed.document_id,
    )
    resolution = PackageResolutionResult(
        documents=[failed],
        document_results=[DocumentExtractionResult(document=failed)],
        candidate_requirements=[],
        reconciled_requirements=[],
        external_references=[],
        unresolved_issues=[issue],
    )

    report = build_compliance_report(resolution)

    assert report.requirements == []
    assert report.summary.documents_failed == 1
    assert report.summary.total_requirements == 0


def test_compliance_workbook_contains_package_matrix_and_supporting_sheets(
    tmp_path: Path,
) -> None:
    """One workbook projects active requirements and all their source evidence."""
    resolution, draft = report_resolution()
    main_document = resolution.documents[0]
    resolution.reconciled_requirements[0].evidence.append(
        Evidence(
            document_id=main_document.document_id,
            document_name=main_document.filename,
            section="Section 2.6",
            page=6,
            text='=HYPERLINK("https://example.invalid","source text")',
        )
    )
    report = build_compliance_report(resolution, ReportModel(draft))
    output = tmp_path / "outputs" / "compliance_matrix.xlsx"

    saved_path = export_compliance_matrix_workbook(report, output)

    assert saved_path == output.resolve()
    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook.sheetnames == [
            "Compliance Matrix",
            "Requirement Details",
            "Issues - Human Review",
            "Document Register",
        ]
        matrix = workbook["Compliance Matrix"]
        assert matrix.freeze_panes == "A5"
        assert matrix.auto_filter.ref == "A4:H6"
        assert [cell.value for cell in matrix[4]] == [
            "Item",
            "Requirement (from RFP)",
            "Sec.",
            "M/R",
            "Resp.",
            "Status",
            "Pg",
            "Comments",
        ]
        assert matrix.max_row == 6
        assert matrix["A5"].value == 1
        assert matrix["A6"].value == 2
        assert matrix["C5"].value == "1"
        assert matrix["G5"].value == 1
        assert matrix["C6"].value == "2"
        assert matrix["G6"].value == 2
        assert [matrix[f"D{row}"].value for row in (5, 6)] == [
            "Mandatory",
            "Required",
        ]
        assert all(matrix[f"E{row}"].value is None for row in (5, 6))
        assert all(matrix[f"F{row}"].value is None for row in (5, 6))
        assert "timing" in matrix["H6"].value.casefold()

        details = workbook["Requirement Details"]
        assert details.max_row == 4
        assert [details[f"B{row}"].value for row in range(2, 5)] == [
            "REQ-0001",
            "REQ-0001",
            "REQ-0002",
        ]
        assert details["E3"].value == main_document.filename
        assert details["H3"].value.startswith("'=")
        assert details["H3"].data_type == "s"

        issues = workbook["Issues - Human Review"]
        assert issues.max_row == 4
        assert "Unclear Timing" in [issues[f"C{row}"].value for row in range(2, 5)]

        register = workbook["Document Register"]
        assert register.max_row == 4
        assert [register[f"B{row}"].value for row in range(2, 5)] == [
            "main_solicitation.pdf",
            "amendment_001.pdf",
            "Annex_E.pdf",
        ]
        assert all(
            cell.data_type != "f"
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
        )
    finally:
        workbook.close()


def test_default_package_outputs_use_requested_paths() -> None:
    """The full runner defaults to one authoritative JSON and one workbook."""
    assert DEFAULT_COMPLIANCE_REPORT_PATH.name == "compliance_report.json"
    assert DEFAULT_COMPLIANCE_MATRIX_PATH.name == "compliance_matrix.xlsx"
    assert DEFAULT_COMPLIANCE_REPORT_PATH.parent.name == "outputs"
    assert (
        DEFAULT_COMPLIANCE_MATRIX_PATH.parent == DEFAULT_COMPLIANCE_REPORT_PATH.parent
    )


def test_full_runner_saves_authoritative_json_and_workbook(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The package runner writes both requested outputs from one final report."""
    resolution, draft = report_resolution()
    report = build_compliance_report(resolution, ReportModel(draft))
    json_path = tmp_path / "outputs" / "compliance_report.json"
    workbook_path = tmp_path / "outputs" / "compliance_matrix.xlsx"
    monkeypatch.setattr(
        main,
        "run_solicitation_package",
        lambda **_kwargs: resolution,
    )
    monkeypatch.setattr(main, "reconcile_package", lambda _extraction: resolution)
    monkeypatch.setattr(main, "resolve_package", lambda _reconciliation: resolution)
    monkeypatch.setattr(main, "build_compliance_report", lambda _resolution: report)
    events: list[tuple[str, str]] = []

    result = main.run_compliance_solicitation_package(
        model=object(),
        output_path=json_path,
        spreadsheet_output_path=workbook_path,
        observe=lambda kind, **fields: events.append(
            (kind, fields.get("stage", fields.get("data", {}).get("code", "")))
        ),
    )

    assert result == report
    assert json_path.is_file()
    assert workbook_path.is_file()
    assert [event for event in events if event[0].startswith("stage.")] == [
        (kind, stage)
        for stage in ("extract", "reconcile", "resolve", "classify", "export")
        for kind in ("stage.started", "stage.completed")
    ]
    assert ("run.warning", "DOCUMENTS_FAILED") in events
    assert json.loads(json_path.read_text(encoding="utf-8")) == report.model_dump(
        mode="json"
    )
